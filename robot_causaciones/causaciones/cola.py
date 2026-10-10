import logging
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from causaciones.modelos import FilaCola

log = logging.getLogger("bot")

ENCABEZADOS = ["IdLista", "Factura", "NIT", "Proveedor", "Valor", "CentroCosto", "FechaFactura"]


def _a_entero(valor: Any) -> int | None:
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, int):
        return valor
    if isinstance(valor, float):
        return int(valor) if valor.is_integer() else None
    texto = str(valor).strip()
    return int(texto) if texto.isascii() and texto.isdigit() else None


def leer_cola(ruta: Path) -> list[FilaCola]:
    """Lee la primera hoja del libro según los encabezados; ignora las filas sin IdLista."""
    libro = load_workbook(ruta, data_only=True)
    filas = libro.worksheets[0].iter_rows(values_only=True)

    columnas: dict[str, int] | None = None
    for fila in filas:
        nombres = ["" if c is None else str(c).strip().casefold() for c in fila]
        if "idlista" in nombres:
            columnas = {n: i for i, n in enumerate(nombres) if n}
            break
    if columnas is None:
        raise ValueError("No se encontró la columna IdLista en la primera hoja del libro")
    faltantes = [h for h in ENCABEZADOS if h.casefold() not in columnas]
    if faltantes:
        raise ValueError(f"Faltan columnas en el libro: {', '.join(faltantes)}")

    def celda(fila: tuple, nombre: str) -> Any:
        i = columnas[nombre.casefold()]
        return fila[i] if i < len(fila) else None

    cola: list[FilaCola] = []
    for fila in filas:
        crudo = celda(fila, "IdLista")
        if crudo is None or str(crudo).strip() == "":
            continue
        id_lista = _a_entero(crudo)
        if id_lista is None:
            log.warning("Fila ignorada: IdLista no es un entero (%r)", crudo)
            continue
        cola.append(
            FilaCola(
                id_lista=id_lista,
                factura=celda(fila, "Factura"),
                nit=celda(fila, "NIT"),
                proveedor=celda(fila, "Proveedor"),
                valor=celda(fila, "Valor"),
                centro=celda(fila, "CentroCosto"),
                fecha=celda(fila, "FechaFactura"),
            )
        )
    return cola
